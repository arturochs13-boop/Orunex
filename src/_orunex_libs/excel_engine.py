"""
excel_engine.py — Orunex Excel Mastery Engine V40.6 (MATRIX MODE)
Motor profesional de generación y edición de hojas de cálculo.
Diseño ejecutivo, analítica multiserie y blindaje de datos.
"""
import os
import json
import uuid
import re
import time

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
#  PALETA EJECUTIVA (V40.6)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
PALETTE = {
    "header_bg":     "1B2A4A",  # Navy profundo
    "header_fg":     "FFFFFF",
    "subheader_bg":  "2C3E6B",  # Navy medio
    "accent":        "C9A84C",  # Gold ejecutivo
    "row_alt":       "EDF1F7",  # Gris azulado claro
    "row_normal":    "FFFFFF",
    "total_bg":      "1B2A4A",
    "total_fg":      "C9A84C",
    "negative":      "CC3333",
    "positive":      "2E7D32",
    "border":        "B0BEC5",
    "highlight":     "FFF3CD",
}

def _get_styles():
    """Genera los estilos ejecutivos reutilizables."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side # type: ignore
    thin = Side(style='thin', color=PALETTE["border"])
    medium = Side(style='medium', color=PALETTE["header_bg"])

    return {
        "header_font": Font(bold=True, color=PALETTE["header_fg"], size=11, name="Calibri"),
        "header_fill": PatternFill(start_color=PALETTE["header_bg"], end_color=PALETTE["header_bg"], fill_type="solid"),
        "header_align": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "header_border": Border(left=thin, right=thin, top=medium, bottom=medium),
        "data_font": Font(size=10, name="Calibri"),
        "data_align": Alignment(vertical="center"),
        "data_border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "alt_fill": PatternFill(start_color=PALETTE["row_alt"], end_color=PALETTE["row_alt"], fill_type="solid"),
        "normal_fill": PatternFill(fill_type=None),
        "total_font": Font(bold=True, color=PALETTE["total_fg"], size=11, name="Calibri"),
        "total_fill": PatternFill(start_color=PALETTE["total_bg"], end_color=PALETTE["total_bg"], fill_type="solid"),
        "total_border": Border(left=thin, right=thin, top=medium, bottom=medium),
        "negative_font": Font(color=PALETTE["negative"], size=10, name="Calibri"),
        "positive_font": Font(color=PALETTE["positive"], size=10, name="Calibri"),
        "pct_format": '0.0%',
        "currency_format": '$#,##0.00',
        "number_format": '#,##0.00',
        "integer_format": '#,##0',
        "date_format": 'DD/MM/YYYY',
    }

def _detect_col_type(header: str, sample_values: list) -> str:
    h = header.lower()
    if any(k in h for k in ["fecha", "date", "día", "mes", "año"]): return "date"
    if any(k in h for k in ["%", "porcentaje", "tasa", "variación", "margen"]): return "percentage"
    if any(k in h for k in ["precio", "costo", "gasto", "ingreso", "pago", "monto", "total", "saldo", "pago"]): return "currency"
    if any(k in h for k in ["cantidad", "qty", "unidades", "num", "número", "numero"]): return "integer"
    
    numeric_count = sum(1 for v in sample_values if isinstance(v, (int, float)))
    if numeric_count > len(sample_values) * 0.6:
        if all(isinstance(v, (int, float)) and 0 <= v <= 1 for v in sample_values if isinstance(v, (int, float))):
            return "percentage"
        return "number"
    return "text"

def _validate_formula(formula: str) -> bool:
    if not formula or not formula.startswith("="): return False
    return formula.count('(') == formula.count(')')

def _apply_cell_format(cell, col_type: str, styles: dict, is_negative: bool = False):
    if col_type == "currency": cell.number_format = styles["currency_format"]
    elif col_type == "percentage": cell.number_format = styles["pct_format"]
    elif col_type == "integer": cell.number_format = styles["integer_format"]
    elif col_type == "number": cell.number_format = styles["number_format"]
    elif col_type == "date": cell.number_format = styles["date_format"]
    if is_negative: cell.font = styles["negative_font"]

def _auto_width(ws):
    from openpyxl.utils import get_column_letter # type: ignore
    for col in ws.columns:
        max_len = 0
        for cell in col:
            try:
                val_len = len(str(cell.value)) if cell.value else 0
                if val_len > max_len: max_len = val_len
            except: pass
        adjusted = min(max_len + 4, 50)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(10, adjusted)

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
#  GENERACIÓN Y MODIFICACIÓN (MATRIX MODE V40.6)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

def modify_excel(instructions: str, input_path: str, out_dir: str) -> str:
    print(f"[EXCEL ENGINE V40.6] Matrix Mode: {input_path}")
    try:
        import openpyxl # type: ignore
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.chart import PieChart, BarChart, LineChart, Reference
    except: return "Error: openpyxl missing."

    if not os.path.exists(input_path): return f"Error: Archivo no hallado: {input_path}"

    # Carga de datos original
    wb_in = openpyxl.load_workbook(input_path, data_only=True)
    ws_in = wb_in.active
    raw_data = []
    for row in ws_in.iter_rows(max_row=1000, values_only=True):
        raw_data.append([v if v is not None else "" for v in row])

    # IA Optimization (Summary)
    from .ollama_client import call_ollama
    data_summary = json.dumps(raw_data[:20], ensure_ascii=False)
    prompt = f"Expert Data Architect: Transform this dataset. Request: {instructions}\nData: {data_summary}\nReturn JSON ONLY: {{'title','headers','analysis','chart_type','formulas'}}"
    
    res = call_ollama("orunex_v37", prompt, options={"temperature": 0.05})
    json_str = _extract_json(res.get("response", "{}"))
    
    try:
        meta = json.loads(json_str)
        title = str(meta.get("title", "Orunex Matrix Report"))[:31]
        headers = meta.get("headers") or []
        formulas = meta.get("formulas") or {}
        chart_type = str(meta.get("chart_type", "")).lower()
    except:
        title, headers, formulas, chart_type = "Matrix Report", [], {}, ""

    # Robust fallback: if LLM returned no headers, use raw_data row 0 or generate generic ones
    if not headers:
        if raw_data and raw_data[0]:
            headers = [str(v) if v else f"Col{i+1}" for i, v in enumerate(raw_data[0])]
        else:
            headers = ["A", "B", "C", "D", "E"]

    rows = raw_data[1:] if len(raw_data) > 1 else raw_data

    # Construcción Workbook Blindado
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = title
    styles = _get_styles()
    all_headers = list(headers)
    for f in formulas: 
        if f not in all_headers: all_headers.append(f)

    # Detect types & Populate Headers
    col_types = {}
    for c_idx, h in enumerate(all_headers, 1):
        samples = [r[c_idx-1] for r in rows[:10] if c_idx-1 < len(r)]
        col_types[c_idx] = _detect_col_type(h, samples)
        cell = ws_out.cell(row=1, column=c_idx, value=h)
        cell.font = styles["header_font"]; cell.fill = styles["header_fill"]
        cell.alignment = styles["header_align"]; cell.border = styles["header_border"]

    # Fill Data & Formulas
    last_row = 1
    for r_idx, r_data in enumerate(rows, 2):
        last_row = r_idx
        is_alt = r_idx % 2 == 0
        for c_idx, val in enumerate(r_data, 1):
            cell = ws_out.cell(row=r_idx, column=c_idx, value=val)
            cell.font = styles["data_font"]; cell.border = styles["data_border"]
            cell.fill = styles["alt_fill"] if is_alt else styles["normal_fill"]
            _apply_cell_format(cell, col_types.get(c_idx, "text"), styles)

        for f_name, f_tpl in formulas.items():
            f_idx = all_headers.index(f_name) + 1
            formula = f_tpl.replace("{row}", str(r_idx))
            if _validate_formula(formula):
                cell = ws_out.cell(row=r_idx, column=f_idx, value=formula)
                cell.font = styles["data_font"]; cell.border = styles["data_border"]
                cell.fill = styles["alt_fill"] if is_alt else styles["normal_fill"]
                _apply_cell_format(cell, col_types.get(f_idx, "number"), styles)

    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    #  MATRIX CHART ENGINE V40.6
    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    chart_keywords = ["gráfica", "graficar", "chart", "pastel", "pie", "barras", "comparar", "analítica"]
    if any(k in instructions.lower() for k in chart_keywords) and last_row > 2:
        try:
            from openpyxl.chart import BarChart, LineChart, PieChart, Reference
            
            data_cols = []
            for c_idx, ct in col_types.items():
                if ct in ("currency", "number", "integer"):
                    h_low = all_headers[c_idx-1].lower()
                    if not any(k in h_low for k in ["id", "código", "índice"]):
                        data_cols.append(c_idx)
            
            if not data_cols: data_cols = [2]
            
            def build_chart():
                if "pie" in chart_type or "pastel" in instructions.lower():
                    c = PieChart()
                    c.title = f"Distribución: {all_headers[data_cols[0]-1]}"
                elif "line" in chart_type or "tendencia" in instructions.lower():
                    c = LineChart(); c.title = "Tendencia Matrix Orunex"
                else:
                    c = BarChart(); c.title = "Análisis Comparativo Orunex"
                    c.type = "col" # Grouped
                
                for d_col in data_cols:
                    d_ref = Reference(ws_out, min_col=d_col, min_row=1, max_row=last_row)
                    c.add_data(d_ref, titles_from_data=True)
                
                l_ref = Reference(ws_out, min_col=1, min_row=2, max_row=last_row)
                c.set_categories(l_ref)
                c.style = 10
                return c

            ws_out.add_chart(build_chart(), "E2")
            ws_dedicated = wb_out.create_sheet("📊 ANALÍTICA MATRIX")
            ws_dedicated.add_chart(build_chart(), "B2")
            ws_dedicated.sheet_view.showGridLines = False
            print(f"[MATRIX] ✅ Gráficas Multiserie inyectadas ({len(data_cols)} series).")
        except Exception as e:
            print(f"[MATRIX ERR] Chart engine: {e}")

    _auto_width(ws_out)
    ws_out.freeze_panes = "B2"
    ws_out.sheet_view.showGridLines = False
    
    path_out = os.path.join(out_dir, f"Orunex_Matrix_{uuid.uuid4().hex[:4]}.xlsx")
    os.makedirs(out_dir, exist_ok=True)
    wb_out.save(path_out)
    return path_out

def generate_excel(instructions: str, out_dir: str) -> str:
    """
    SOVEREIGN EXCEL GENERATOR V51.0
    Creates a fully populated professional Excel from scratch using the GGUF model.
    Does NOT delegate to modify_excel (which requires existing data).
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from .ollama_client import call_ollama
    print(f"[EXCEL ENGINE V51] Generating new Excel: {instructions[:60]}")

    # Ask GGUF to generate structured data
    prompt = (
        f"You are an expert data analyst. Create a professional Excel dataset about: {instructions}\n"
        "Return ONLY a JSON object with these fields:\n"
        '{"title": "Sheet title", "headers": ["Col1", "Col2", ...], '
        '"rows": [[val1, val2, ...], [val1, val2, ...]], "summary": "Brief description"}\n'
        "Include 8-15 realistic data rows. No markdown, no explanation."
    )
    res = call_ollama("orunex_v37", prompt, options={"temperature": 0.1, "num_predict": 1500})
    raw = res.get("response", "{}")

    # Extract JSON
    import json
    json_str = _extract_json(raw)
    try:
        meta = json.loads(json_str)
        title = str(meta.get("title", "Orunex Data"))[:31]
        headers = meta.get("headers") or ["Item", "Valor", "Descripcion"]
        rows = meta.get("rows") or []
    except:
        title = "Orunex Matrix"
        headers = ["Concepto", "Valor", "Porcentaje", "Estado"]
        rows = [
            ["Dato 1", 1250.0, 0.25, "Activo"],
            ["Dato 2", 980.0, 0.20, "Activo"],
            ["Dato 3", 2100.0, 0.42, "Critico"],
            ["Dato 4", 640.0, 0.13, "Pendiente"],
        ]

    # Build workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    styles = _get_styles()
    ws.sheet_view.showGridLines = False

    # Header row
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=str(h))
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["header_align"]
        cell.border = styles["header_border"]

    # Data rows
    for r_idx, row_data in enumerate(rows, 2):
        is_alt = r_idx % 2 == 0
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = styles["data_font"]
            cell.border = styles["data_border"]
            cell.fill = styles["alt_fill"] if is_alt else styles["normal_fill"]
            if c_idx <= len(headers):
                col_types = _detect_col_type(str(headers[c_idx-1]), [val])
                _apply_cell_format(cell, col_types, styles)

    _auto_width(ws)
    ws.freeze_panes = "A2"

    path_out = os.path.join(out_dir, f"Orunex_Matrix_{uuid.uuid4().hex[:4]}.xlsx")
    os.makedirs(out_dir, exist_ok=True)
    wb.save(path_out)
    print(f"[EXCEL ENGINE V51] Saved: {path_out}")
    return path_out

def _extract_json(raw: str) -> str:
    try:
        raw = re.sub(r"```json|```", "", raw)
        start = raw.find("{")
        if start == -1: return "{}"
        content = raw[start:].strip()
        # Auto-repair braces
        o_br, c_br = content.count('{'), content.count('}')
        if o_br > c_br: content += '}' * (o_br - c_br)
        last_brace = content.rfind('}')
        if last_brace != -1: content = content[:last_brace+1]
        return content
    except: return "{}"
