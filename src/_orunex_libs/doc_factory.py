import os
import json
import uuid
import re

from .headless_web import search_and_summarize
from .ollama_client import call_ollama

def clean_json_resp(ai_resp):
    if "</think>" in ai_resp: ai_resp = ai_resp.split("</think>")[-1].strip()
    if "```json" in ai_resp:
        return ai_resp.split("```json")[1].split("```")[0].strip()
    elif "```" in ai_resp:
        return ai_resp.split("```")[1].split("```")[0].strip()
    else:
        start = ai_resp.find("[")
        if start == -1: start = ai_resp.find("{")
        end = ai_resp.rfind("]")
        if end == -1: end = ai_resp.rfind("}")
        if start != -1 and end != -1:
            return ai_resp[start:end+1]
    return "[]"

def generate_excel(instructions, out_dir):
    """Delegado al Excel Mastery Engine V37.2"""
    from .excel_engine import generate_excel as _engine_gen
    return _engine_gen(instructions, out_dir)


def _generate_excel_legacy(instructions, out_dir):
    """[LEGACY] Excel Architect V37.1 — preservado como referencia."""
    print(f"[DOC_FACTORY] Excel Architect V37.1: '{instructions}'")
    try:
        import openpyxl  # type: ignore
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers  # type: ignore
        from openpyxl.utils import get_column_letter  # type: ignore
    except:
        print("[ERR] openpyxl no instalado")
        return None

    prompt = f"""Eres un ARQUITECTO DE DATOS DE ÉLITE. Genera una hoja de cálculo Excel PERFECTA.

SOLICITUD: '{instructions}'

REGLAS ABSOLUTAS:
1. Genera MÍNIMO 15 FILAS de datos realistas, variados y coherentes. NO repitas valores.
2. Los encabezados deben ser profesionales y descriptivos.
3. Incluye columnas calculadas con FÓRMULAS Excel válidas. Usa '{{row}}' como placeholder de fila.
4. Si hay montos, usa números reales con decimales. Si hay porcentajes, usa valores entre 0 y 1.
5. Incluye una FILA FINAL de TOTALES con fórmulas =SUM().

FORMATO JSON OBLIGATORIO (ejemplo):
{{
  "title": "Presupuesto Departamental 2026",
  "headers": ["Departamento", "Gastos Fijos", "Gastos Variables", "Total", "% del Total"],
  "rows": [
    ["Desarrollo", 45000.00, 12500.00],
    ["Marketing", 28000.00, 35000.00],
    ["RRHH", 22000.00, 8000.00]
  ],
  "formulas": {{
    "Total": "=B{{row}}+C{{row}}",
    "% del Total": "=D{{row}}/SUM(D2:D16)"
  }},
  "totals_row": ["TOTAL", "=SUM(B2:B16)", "=SUM(C2:C16)", "=SUM(D2:D16)", ""]
}}

Devuelve ÚNICAMENTE el JSON. Sin texto adicional."""

    res = call_ollama("orunex_v37", prompt, options={"temperature": 0.1})
    raw_resp = res.get("response", "{}")
    json_str = clean_json_resp_obj(raw_resp)

    try:
        data = json.loads(json_str)
        title    = data.get("title", "Hoja Orunex")
        headers  = data.get("headers", [])
        rows     = data.get("rows", [])
        formulas = data.get("formulas", {})
        totals   = data.get("totals_row", [])
        if not headers or not rows:
            raise ValueError("Datos vacíos del LLM")
    except Exception as e:
        print(f"[WARN] LLM JSON fallido: {e}. Usando fallback profesional.")
        title = "Presupuesto Empresarial 2026"
        headers = ["Departamento", "Gastos Fijos", "Gastos Variables", "Total", "% Presupuesto"]
        rows = [
            ["Desarrollo", 45000, 12500],
            ["Marketing", 28000, 35000],
            ["Recursos Humanos", 22000, 8000],
            ["Infraestructura", 38000, 15000],
            ["Legal", 18000, 5000],
            ["Ventas", 32000, 22000],
            ["Soporte Técnico", 25000, 9000],
            ["Investigación", 41000, 18000],
            ["Operaciones", 29000, 11000],
            ["Finanzas", 20000, 7000],
        ]
        formulas = {"Total": "=B{row}+C{row}", "% Presupuesto": "=D{row}/SUM(D2:D11)"}
        totals = ["TOTAL", "=SUM(B2:B11)", "=SUM(C2:C11)", "=SUM(D2:D11)", ""]

    # Build column map
    all_headers = list(headers)
    for fcol in formulas:
        if fcol not in all_headers:
            all_headers.append(fcol)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]  # Excel limit: 31 chars

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    alt_fill    = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    total_font  = Font(bold=True, size=11)
    total_fill  = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    red_font    = Font(color="CC0000")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # Write headers
    for c_idx, h in enumerate(all_headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Write data rows + formulas
    for r_idx, row_data in enumerate(rows, 2):
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            if r_idx % 2 == 0:
                cell.fill = alt_fill
            # Format numbers
            if isinstance(val, (int, float)):
                if val < 0:
                    cell.font = red_font
                if abs(val) >= 100:
                    cell.number_format = '#,##0.00'

        # Write formula columns
        for fcol, ftemplate in formulas.items():
            if fcol in all_headers:
                c_idx = all_headers.index(fcol) + 1
                formula = ftemplate.replace("{row}", str(r_idx))
                cell = ws.cell(row=r_idx, column=c_idx, value=formula)
                cell.border = thin_border
                if r_idx % 2 == 0:
                    cell.fill = alt_fill
                # Format percentage columns
                if "%" in fcol:
                    cell.number_format = '0.0%'
                else:
                    cell.number_format = '#,##0.00'

    # Write totals row
    if totals:
        total_row_idx = len(rows) + 2
        for c_idx, val in enumerate(totals, 1):
            cell = ws.cell(row=total_row_idx, column=c_idx, value=val)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = thin_border
            if isinstance(val, str) and val.startswith("="):
                cell.number_format = '#,##0.00'

    # Auto-size columns
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 5, 45)

    # Freeze top row
    ws.freeze_panes = "A2"

    safe_name = "".join([c for c in instructions[:25].replace(' ', '_') if c.isalnum() or c == '_'])
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"{safe_name}_{uuid.uuid4().hex[:4]}.xlsx")
    wb.save(path)
    print(f"[DOC_FACTORY] Excel generado: {path}")
    return path


def clean_json_resp_obj(ai_resp: str) -> str:
    """NUCLEAR JSON EXTRACTOR V22.0 - Destruye razonamientos de R1."""
    if not ai_resp: return "{}"
    if "</think>" in ai_resp:
        ai_resp = ai_resp.split("</think>")[-1].strip()
    
    # Intento 1: Markdown blocks
    for marker in ["```json", "```"]:
        if marker in ai_resp:
            parts = ai_resp.split(marker)
            if len(parts) > 1:
                content = parts[1].split("```")[0].strip()
                if content: return content
                
    # Intento 2: Busqueda de limites estructurales
    for start_ch, end_ch in ( ('{', '}'), ('[', ']') ):
        s = ai_resp.find(start_ch)
        e = ai_resp.rfind(end_ch)
        if s != -1 and e != -1 and e > s:
            content = ai_resp[s:e+1]
            if content: return content
            
    return ai_resp.strip()


def generate_word(instructions, out_dir):
    print(f"[DOC_FACTORY] Generando Word: {instructions}")
    try:
        import docx # type: ignore
    except: return None
    
    prompt = f"Escribe el contenido para un documento formal sobre: '{instructions}'. \n\nDevuelve ESTRICTAMENTE un JSON con formato de lista de diccionarios. El output solo debe ser el JSON puro. Sin saludos ni marcas de agua.\n[\n  {{\"type\": \"title\", \"text\": \"Titulo Principal\"}},\n  {{\"type\": \"paragraph\", \"text\": \"Un parrafo largo...\"}},\n  {{\"type\": \"heading\", \"text\": \"Subtitulo\"}},\n  {{\"type\": \"bullet\", \"text\": \"Un punto clave\"}}\n]"
    res = call_ollama("orunex-r1", prompt, options={"temperature": 0.2})
    json_str = clean_json_resp(res.get("response", "[]"))
    
    try:
        data = json.loads(json_str)
        if not data or not isinstance(data, list): raise ValueError("Format error")
    except:
        data = [
            {"type": "title", "text": "Documento Base"},
            {"type": "paragraph", "text": f"Contenido auto-generado sobre: {instructions}"}
        ]

    doc = docx.Document()
    
    for item in data:
        t = item.get("type", "paragraph")
        text = str(item.get("text", "")).strip()
        if not text: continue
        
        if t == "title":
            doc.add_heading(text, level=0)
        elif t == "heading":
            doc.add_heading(text, level=1)
        elif t == "bullet":
            doc.add_paragraph(text, style='List Bullet')
        else:
            doc.add_paragraph(text)

    safe_name = "".join([c for c in instructions[:20].replace(' ', '_') if c.isalnum() or c == '_'])
    path = os.path.join(out_dir, f"Orunex_Doc_{safe_name}_{uuid.uuid4().hex[:4]}.docx")
    doc.save(path)
    return path

def generate_pdf(instructions, out_dir):
    print(f"[DOC_FACTORY] Generando PDF: {instructions}")
    try:
        from fpdf import FPDF # type: ignore
    except:
        return None

    def _safe(text):
        """Sanitize string to Latin-1 for fpdf compatibility and prevent layout breaks."""
        text = text.encode('latin-1', errors='replace').decode('latin-1')
        # Prevent "Not enough horizontal space" by breaking extremely long words/lines
        words = text.split(" ")
        new_words = []
        for w in words:
            if len(w) > 40: # Max word length to avoid overflow
                new_words.append(w[:40] + "...")
            else:
                new_words.append(w)
        return " ".join(new_words)

    prompt = (
        f"Escribe el texto estructurado para un PDF ejecutivo sobre: '{instructions}'.\n\n"
        "Usa '#' para titulos, '##' para subtitulos y texto normal para parrafos.\n"
        "Genera minimo 5 secciones con contenido detallado y profesional.\n"
        "NO uses caracteres especiales como emojis ni simbolos Unicode raros.\n"
        "NO generes lineas de guiones o puntos extremadamente largas (ej: --------).\n"
        "Solo texto plano, sin saludos ni firmas."
    )
    res = call_ollama("orunex-r1", prompt, options={"temperature": 0.2})
    raw = res.get("response", f"# Documento PDF\n\nContenido sobre: {instructions}")
    if "</think>" in raw:
        raw = raw.split("</think>")[-1].strip()

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    # Calculate safe width
    effective_width = pdf.w - 2 * pdf.l_margin

    for line in raw.split("\n"):
        line = line.strip()
        if not line:
            pdf.ln(5)
            continue
        
        try:
            if line.startswith("# "):
                pdf.set_font("helvetica", "B", 20)
                pdf.multi_cell(effective_width, 10, txt=_safe(line[2:]), align="L")
                pdf.set_font("helvetica", "", 12)
            elif line.startswith("## "):
                pdf.set_font("helvetica", "B", 16)
                pdf.multi_cell(effective_width, 10, txt=_safe(line[3:]), align="L")
                pdf.set_font("helvetica", "", 12)
            elif line.startswith("### "):
                pdf.set_font("helvetica", "B", 14)
                pdf.multi_cell(effective_width, 10, txt=_safe(line[4:]), align="L")
                pdf.set_font("helvetica", "", 12)
            elif line.startswith("- ") or line.startswith("* "):
                pdf.set_font("helvetica", "", 12)
                pdf.multi_cell(effective_width, 8, txt=_safe("  - " + line[2:]), align="L")
            else:
                pdf.set_font("helvetica", "", 12)
                pdf.multi_cell(effective_width, 8, txt=_safe(line), align="L")
        except:
            continue # Skip problematic lines to avoid full crash

    safe_name = "".join([c for c in instructions[:20].replace(' ', '_') if c.isalnum() or c == '_'])
    path = os.path.join(out_dir, f"Orunex_Doc_{safe_name}_{uuid.uuid4().hex[:4]}.pdf")
    pdf.output(path)
    return path

def modify_excel(instructions, input_path, out_dir):
    """Delegado al Excel Mastery Engine V37.2"""
    from .excel_engine import modify_excel as _engine_mod
    return _engine_mod(instructions, input_path, out_dir)


def _modify_excel_legacy(instructions, input_path, out_dir):
    """[LEGACY] Modificador Excel básico — preservado como referencia."""
    print(f"[DOC_FACTORY] Modificando Excel: {input_path}")
    try:
        import openpyxl  # type: ignore
        from openpyxl.styles import Font, PatternFill, Alignment  # type: ignore
        from openpyxl.utils import get_column_letter  # type: ignore
    except:
        return f"Error: openpyxl no instalado."

    if not os.path.exists(input_path):
        return f"Error: Archivo no encontrado en {input_path}"

    # 1. Leer datos existentes (Max 150 filas para proteger contexto)
    try:
        wb_in = openpyxl.load_workbook(input_path, data_only=True)
        ws_in = wb_in.active
        
        raw_data = []
        for row in ws_in.iter_rows(max_row=150, values_only=True): # type: ignore
            raw_data.append(list(row))
        
        # Convertir a texto representativo
        data_summary = json.dumps(raw_data, ensure_ascii=False)
    except Exception as e:
        return f"Error leyendo Excel: {e}"

    prompt = f"""Eres un experto en analisis y diseño de Excel. 

DATOS ORIGINALES (Primeras 150 filas):
{data_summary}

SOLICITUD DEL USUARIO: "{instructions}"

INSTRUCCIONES:
1. Analiza los datos y reorganizalos, mejoralos o filtra segun pida el usuario.
2. Devuelve ESTRICTAMENTE un JSON con este formato:
{{
  "headers": ["Col1", "Col2", ...],
  "rows": [ ["val1", "val2", ...], ... ],
  "styles": {{
     "theme": "dark" | "blue" | "modern" | "clean",
     "highlight_cols": ["NombreColA", "NombreColB"]
  }}
}}

No hables, solo JSON puro."""

    res = call_ollama("orunex-r1", prompt, options={"temperature": 0.1})
    raw_resp = res.get("response", "{{}}")
    json_str = clean_json_resp_obj(raw_resp)

    try:
        data = json.loads(json_str)
        headers = data.get("headers", [])
        rows = data.get("rows", [])
        style = data.get("styles", {})
    except:
        # Fallback si falla el JSON
        headers = raw_data[0] if raw_data else ["Error"]
        rows = raw_data[1:] if len(raw_data) > 1 else [["Fallo el procesamiento IA"]]
        style = {}

    # 2. Crear nuevo libro con los cambios
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Orunex Smart Edit" # type: ignore

    # Definir paletas segun el 'theme'
    theme_name = style.get("theme", "blue").lower()
    
    colors = {
        "dark":  {"bg": "212121", "fg": "FFFFFF", "alt": "333333"},
        "blue":  {"bg": "2E75B6", "fg": "FFFFFF", "alt": "DCE6F1"},
        "modern": {"bg": "7030A0", "fg": "FFFFFF", "alt": "F2DCF7"},
        "clean": {"bg": "BFBFBF", "fg": "000000", "alt": "F2F2F2"}
    }
    
    pal = colors.get(theme_name, colors["blue"])
    header_fill = PatternFill(start_color=pal["bg"], end_color=pal["bg"], fill_type="solid")
    header_font = Font(bold=True, color=pal["fg"])
    alt_fill = PatternFill(start_color=pal["alt"], end_color=pal["alt"], fill_type="solid")

    # Escribir Headers
    for c_idx, h in enumerate(headers, 1):
        cell = ws_out.cell(row=1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Escribir Filas
    for r_idx, row_val in enumerate(rows, 2):
        for c_idx, val in enumerate(row_val, 1):
            cell = ws_out.cell(row=r_idx, column=c_idx, value=val)
            if r_idx % 2 == 0:
                cell.fill = alt_fill

    # Auto-ajuste
    for col in ws_out.columns: # type: ignore
        max_len = 0
        column = col[0].column_letter # type: ignore
        for cell in col:
            try:
                if len(str(cell.value)) > max_len:
                    max_len = len(str(cell.value))
            except: pass
        ws_out.column_dimensions[column].width = min(max_len + 4, 50)

    filename = f"SmartEdit_{os.path.basename(input_path)}"
    if not filename.endswith(".xlsx"): filename += ".xlsx"
    
    path_out = os.path.join(out_dir, f"{uuid.uuid4().hex[:4]}_{filename}")
    wb_out.save(path_out)
    return path_out

def generate_landing(instructions, out_dir):
    """Orunex V42.0: Premium Agency Web Engineering (Glassmorphism & High-FID)."""
    print(f"[DOC_FACTORY] Generando Landing Page PREMIUM: {instructions}")
    
    # 💎 Elite Design System Prompt (V42.1)
    prompt = f"""Eres un DIRECTOR CREATIVO DE UNA AGENCIA DE MARKETING DE ÉLITE (como IDEO o Frog Design).
    TAREA: Diseña una Landing Page ULTRA-PREMIUM sobre: '{instructions}'.
    
    ESTRUCTURA OBLIGATORIA (V42.0):
    1. NAV: Logo minimalista y botones con efecto hover sutil.
    2. HERO: Titular MASIVO con gradiente de color, subtitulo elegante y boton 'Glass' con pulso.
    3. LOGOS: Fila de logos de clientes (placeholders) para generar confianza.
    4. SERVICES: Grid de tarjetas Glassmorphism (blur, transparent borders).
    5. CTA FINAL: Seccion con fondo de gradiente animado.
    
    REQUISITOS TÉCNICOS:
    - Framework: Tailwind CSS (<script src="https://cdn.tailwindcss.com"></script>).
    - Paleta: Deep Navy (#050510) como base, Cyber Blue y Electric Purple como acentos.
    - Imagenes: Usa <img src="https://images.unsplash.com/photo-NNN?auto=format&fit=crop&w=800&q=80"> (usa IDs reales de Unsplash para Marketing).
    - Efectos: backdrop-blur-xl, rounded-3xl, shadow-2xl, transition-all duration-500.
    
    COMPORTAMIENTO: Devuelve SOLO el HTML completo. Debe ser una obra maestra visual lista para publicar.
    """
    
    res = call_ollama("orunex-r1", prompt, options={"temperature": 0.5})
    raw = res.get("response", "")
    
    # Anti-Think Logic
    if "</think>" in raw: raw = raw.split("</think>")[-1].strip()
    
    # 1. Surgical HTML Extraction
    match = re.search(r'```(?:html)?(.*?)```', raw, re.DOTALL | re.IGNORECASE)
    html = match.group(1).strip() if match else raw.strip()
    
    # 2. Advanced Integrity Check
    if len(html) < 500 or "tailwindcss" not in html.lower():
        print("[WARN] IA fallo en estetica premium. Aplicando 'Executive Template' de Orunex.")
        html = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{instructions}</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;700;800&display=swap" rel="stylesheet">
    <style>
        body {{ background-color: #030303; font-family: 'Plus Jakarta Sans', sans-serif; overflow-x: hidden; }}
        .glass {{ background: rgba(255, 255, 255, 0.03); backdrop-filter: blur(20px); border: 1px solid rgba(255, 255, 255, 0.1); }}
        .grad-text {{ background: linear-gradient(90deg, #00F0FF, #7000FF); -webkit-background-clip: text; -webkit-text-fill-color: transparent; }}
        .blob {{ position: absolute; width: 500px; height: 500px; background: radial-gradient(circle, rgba(0,240,255,0.15) 0%, rgba(0,0,0,0) 70%); z-index: -1; }}
    </style>
</head>
<body class="text-white">
    <div class="blob -top-20 -left-20"></div>
    <div class="blob bottom-0 right-0" style="background: radial-gradient(circle, rgba(112,0,255,0.1) 0%, rgba(0,0,0,0) 70%);"></div>

    <nav class="p-8 flex justify-between items-center max-w-7xl mx-auto">
        <div class="text-2xl font-black tracking-tighter">ORUNEX <span class="grad-text">LABS</span></div>
        <div class="hidden md:flex gap-8 text-sm font-bold uppercase tracking-widest text-gray-400">
            <a href="#" class="hover:text-cyan-400 transition">Agencia</a>
            <a href="#" class="hover:text-cyan-400 transition">Resultados</a>
            <a href="#" class="hover:text-cyan-400 transition">Estrategia</a>
        </div>
        <div class="px-6 py-2 border border-white/20 rounded-full text-sm font-bold hover:bg-white/5 cursor-pointer transition">Contacto</div>
    </nav>

    <header class="py-24 px-8 text-center max-w-5xl mx-auto">
        <div class="inline-block px-4 py-1 bg-white/5 rounded-full text-xs font-bold border border-white/10 mb-8 text-cyan-400 tracking-widest uppercase">
            Soberanía en Marketing V42.0
        </div>
        <h1 class="text-6xl md:text-8xl font-black mb-8 leading-[1.1] tracking-tight">
            Elevamos tu <span class="grad-text">Visión Digital</span> a otro nivel.
        </h1>
        <p class="text-gray-400 text-xl md:text-2xl mb-12 max-w-3xl mx-auto font-medium">
            Diseñamos estrategias disruptivas para la era de la inteligencia artificail. Landing pages que no solo informan, sino que convierten.
        </p>
        <div class="flex flex-col md:flex-row gap-4 justify-center items-center">
            <div class="px-10 py-5 bg-gradient-to-r from-cyan-500 to-blue-600 rounded-2xl font-bold text-lg hover:scale-105 transition transform shadow-[0_0_40px_rgba(6,182,212,0.3)]">
                Empezar Proyecto
            </div>
            <div class="px-10 py-5 glass rounded-2xl font-bold text-lg hover:bg-white/10 transition">
                Saber más
            </div>
        </div>
    </header>

    <section class="max-w-7xl mx-auto px-8 py-20">
        <div class="grid md:grid-cols-3 gap-8">
            <div class="glass p-10 rounded-[2.5rem] hover:border-cyan-500/50 transition duration-500 group">
                <div class="w-16 h-16 bg-cyan-500/20 rounded-2xl mb-8 flex items-center justify-center text-cyan-400 group-hover:scale-110 transition">
                    <svg class="w-8 h-8" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M13 10V3L4 14h7v7l9-11h-7z"/></svg>
                </div>
                <h3 class="text-2xl font-bold mb-4">Optimización Real</h3>
                <p class="text-gray-400 leading-relaxed">Algoritmos avanzados para maximizar el ROI en cada campaña digital.</p>
            </div>
            <div class="glass p-10 rounded-[2.5rem] hover:border-purple-500/50 transition duration-500 group">
                <div class="w-16 h-16 bg-purple-500/20 rounded-2xl mb-8 flex items-center justify-center text-purple-400 group-hover:scale-110 transition">
                    <svg class="w-8 h-8" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M9 19v-6a2 2 0 00-2-2H5a2 2 0 00-2 2v6a2 2 0 002 2h2a2 2 0 002-2zm0 0V9a2 2 0 012-2h2a2 2 0 012 2v10m-6 0a2 2 0 002 2h2a2 2 0 002-2m0 0V5a2 2 0 012-2h2a2 2 0 012 2v14a2 2 0 01-2 2h-2a2 2 0 01-2-2z"/></svg>
                </div>
                <h3 class="text-2xl font-bold mb-4">Análisis Predictivo</h3>
                <p class="text-gray-400 leading-relaxed">No adivinamos, sabemos lo que tu audiencia busca antes de que ellos lo sepan.</p>
            </div>
            <div class="glass p-10 rounded-[2.5rem] hover:border-pink-500/50 transition duration-500 group">
                <div class="w-16 h-16 bg-pink-500/20 rounded-2xl mb-8 flex items-center justify-center text-pink-400 group-hover:scale-110 transition">
                    <svg class="w-8 h-8" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M12 11c0 3.517-1.009 6.799-2.753 9.571m-3.44-2.04l.054-.09a13.916 13.916 0 002.103-4.735m.983-3.172a13.953 13.953 0 01-2.103 4.735m2.235-4.634a13.93 13.93 0 011.832 3.018m1.912-4.176a13.947 13.947 0 01-1.832 3.018m-3.383-1.103a13.933 13.933 0 012.128-4.754m-1.214 4.872a13.957 13.957 0 012.232-4.426m-.054-.11a13.945 13.945 0 00-2.232 4.426m2.232-4.426a13.95 13.95 0 012.126 4.754m-1.215-4.872a13.958 13.958 0 011.215 4.872m-1.215-4.872a13.958 13.958 0 00-1.215 4.872" /></svg>
                </div>
                <h3 class="text-2xl font-bold mb-4">Identidad de Marca</h3>
                <p class="text-gray-400 leading-relaxed">Creamos huellas digitales imborrables en el mercado global.</p>
            </div>
        </div>
    </section>

</body>
</html>"""

    # 3. Persistence (Hardened)
    safe_name = "".join([c for c in instructions[:20].replace(' ', '_') if c.isalnum() or c == '_'])
    fname = "index.html" if "Orunex_Workspace" in out_dir else f"Landing_{safe_name}.html"
    path = os.path.join(out_dir, fname)
    
    try:
        with open(path, "w", encoding="utf-8") as f:
            f.write(html)
        return path
    except Exception as e:
        print(f"[ERR] Error guardando landing page: {e}")
        return None
